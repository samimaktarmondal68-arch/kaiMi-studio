"""Enhanced export service supporting TXT, Markdown, DOCX, PDF, JSON, ZIP."""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

from core.project_manager import ProjectManager
from core.logger import get_logger


class ExportService:

    FORMATS = ["txt", "markdown", "json", "zip", "docx", "pdf"]

    def __init__(self, project_manager: ProjectManager | None = None):
        self.project_manager = project_manager or ProjectManager()
        self._log = get_logger()

    def _load_stage_data(self, project_name, stage):
        project_path = self.project_manager.PROJECTS_DIR / project_name
        if stage == "Research":
            fp = project_path / "research.json"
        elif stage == "Script":
            fp = project_path / "script.json"
        elif stage == "Storyboard":
            fp = project_path / "storyboard.json"
        elif stage == "Image Prompts":
            fp = project_path / "image_prompts.json"
        else:
            return {}
        if not fp.exists():
            return {}
        try:
            with open(fp, "r", encoding="utf-8") as f:
                return json.load(f)
        except (OSError, json.JSONDecodeError):
            return {}

    def export_project(self, project: dict, fmt: str = "zip") -> Path:
        project_name = project.get("name") or "project"
        self._log.info("Export", f"Exporting '{project_name}' as {fmt}")
        export_root = self.project_manager.PROJECTS_DIR / project_name / "exports" / project_name
        export_root.mkdir(parents=True, exist_ok=True)

        research_data = self._load_stage_data(project_name, "Research")
        script_data = self._load_stage_data(project_name, "Script")
        storyboard_data = self._load_stage_data(project_name, "Storyboard")
        image_prompt_data = self._load_stage_data(project_name, "Image Prompts")

        all_data = {
            "project": {k: v for k, v in project.items() if k != "path"},
            "research": research_data,
            "script": script_data,
            "storyboard": storyboard_data,
            "image_prompts": image_prompt_data,
        }

        if fmt == "json":
            return self._export_json(export_root, all_data, project_name)
        elif fmt == "txt":
            return self._export_txt(export_root, all_data, project_name)
        elif fmt == "markdown":
            return self._export_markdown(export_root, all_data, project_name)
        elif fmt == "docx":
            return self._export_docx(export_root, all_data, project_name)
        elif fmt == "pdf":
            return self._export_pdf(export_root, all_data, project_name)
        elif fmt == "zip":
            return self._export_zip(export_root, all_data, project_name)
        else:
            return self._export_zip(export_root, all_data, project_name)

    def _export_json(self, export_root, data, name):
        out = export_root / f"{name}.json"
        with open(out, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=4)
        return out

    def _export_txt(self, export_root, data, name):
        lines = []
        proj = data.get("project", {})
        lines.append(f"Project: {proj.get('name', '')}")
        lines.append(f"Topic: {proj.get('topic', '')}")
        lines.append(f"Status: {proj.get('status', '')}")
        lines.append(f"Created: {proj.get('created', '')}")
        lines.append("")

        research = data.get("research", {})
        if research.get("generated_research"):
            lines.append("=" * 60)
            lines.append("RESEARCH")
            lines.append("=" * 60)
            lines.append(research["generated_research"])
            lines.append("")

        script = data.get("script", {})
        if script.get("script_output"):
            lines.append("=" * 60)
            lines.append("SCRIPT")
            lines.append("=" * 60)
            lines.append(script["script_output"])
            lines.append("")

        storyboard = data.get("storyboard", {})
        for scene in storyboard.get("scenes", []):
            lines.append(f"--- Scene {scene.get('scene_number', '')} ---")
            lines.append(f"Timestamp: {scene.get('timestamp', '')}")
            lines.append(f"Narration: {scene.get('narration', '')}")
            lines.append(f"Visual: {scene.get('visual_description', '')}")
            lines.append(f"Camera: {scene.get('camera_direction', '')}")
            lines.append("")

        prompts = data.get("image_prompts", {})
        for prompt in prompts.get("prompts", []):
            lines.append(f"--- Scene {prompt.get('scene_number', '')} Prompt ---")
            lines.append(prompt.get("full_image_prompt", ""))
            lines.append("")

        out = export_root / f"{name}.txt"
        out.write_text("\n".join(lines), encoding="utf-8")
        return out

    def _export_markdown(self, export_root, data, name):
        lines = []
        proj = data.get("project", {})
        lines.append(f"# {proj.get('name', '')}\n")
        lines.append(f"**Topic:** {proj.get('topic', '')}  ")
        lines.append(f"**Status:** {proj.get('status', '')}  ")
        lines.append(f"**Created:** {proj.get('created', '')}\n")

        research = data.get("research", {})
        if research.get("generated_research"):
            lines.append("## Research\n")
            lines.append(research["generated_research"])
            lines.append("")

        script = data.get("script", {})
        if script.get("script_output"):
            lines.append("## Script\n")
            lines.append(script["script_output"])
            lines.append("")

        storyboard = data.get("storyboard", {})
        if storyboard.get("scenes"):
            lines.append("## Storyboard\n")
            for scene in storyboard["scenes"]:
                lines.append(f"### Scene {scene.get('scene_number', '')}\n")
                lines.append(f"**Timestamp:** {scene.get('timestamp', '')}\n")
                lines.append(f"**Narration:** {scene.get('narration', '')}\n")
                lines.append(f"**Visual:** {scene.get('visual_description', '')}\n")
                lines.append(f"**Camera:** {scene.get('camera_direction', '')}\n")

        prompts = data.get("image_prompts", {})
        if prompts.get("prompts"):
            lines.append("## Image Prompts\n")
            for prompt in prompts["prompts"]:
                lines.append(f"### Scene {prompt.get('scene_number', '')}\n")
                lines.append(f"**{prompt.get('prompt_title', '')}**\n")
                lines.append(f"```\n{prompt.get('full_image_prompt', '')}\n```\n")

        out = export_root / f"{name}.md"
        out.write_text("\n".join(lines), encoding="utf-8")
        return out

    def _export_zip(self, export_root, data, name):
        json_out = self._export_json(export_root, data, name)
        txt_out = self._export_txt(export_root, data, name)
        md_out = self._export_markdown(export_root, data, name)

        zip_path = export_root / f"{name}.zip"
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.write(json_out, f"{name}/{name}.json")
            zf.write(txt_out, f"{name}/{name}.txt")
            zf.write(md_out, f"{name}/{name}.md")

            project_path = self.project_manager.PROJECTS_DIR / name
            for ext in ["research.json", "script.json", "storyboard.json", "image_prompts.json"]:
                src = project_path / ext
                if src.exists():
                    zf.write(src, f"{name}/{ext}")

        return zip_path

    def export_stage(self, project_name: str, stage: str, fmt: str = "txt") -> Path | None:
        data = self._load_stage_data(project_name, stage)
        if not data:
            return None

        export_root = self.project_manager.PROJECTS_DIR / project_name / "exports" / project_name
        export_root.mkdir(parents=True, exist_ok=True)

        filename = f"{stage.lower().replace(' ', '_')}.{fmt}"
        out = export_root / filename

        if fmt == "json":
            with open(out, "w", encoding="utf-8") as f:
                json.dump(data, f, indent=4)
        elif fmt == "txt":
            out.write_text(self._stage_to_txt(data, stage), encoding="utf-8")
        elif fmt == "markdown":
            out.write_text(self._stage_to_md(data, stage), encoding="utf-8")

        return out

    def _stage_to_txt(self, data, stage):
        if stage == "Research":
            return data.get("generated_research", "")
        elif stage == "Script":
            return data.get("script_output", "")
        elif stage == "Storyboard":
            lines = []
            for s in data.get("scenes", []):
                lines.append(f"Scene {s.get('scene_number', '')}: {s.get('narration', '')}")
            return "\n".join(lines)
        elif stage == "Image Prompts":
            lines = []
            for p in data.get("prompts", []):
                lines.append(f"Scene {p.get('scene_number', '')}: {p.get('full_image_prompt', '')}")
            return "\n".join(lines)
        return ""

    def _stage_to_md(self, data, stage):
        if stage == "Research":
            return f"# Research\n\n{data.get('generated_research', '')}"
        elif stage == "Script":
            return f"# Script\n\n{data.get('script_output', '')}"
        elif stage == "Storyboard":
            lines = ["# Storyboard\n"]
            for s in data.get("scenes", []):
                lines.append(f"## Scene {s.get('scene_number', '')}\n")
                lines.append(f"**Narration:** {s.get('narration', '')}\n")
                lines.append(f"**Visual:** {s.get('visual_description', '')}\n")
            return "\n".join(lines)
        elif stage == "Image Prompts":
            lines = ["# Image Prompts\n"]
            for p in data.get("prompts", []):
                lines.append(f"## Scene {p.get('scene_number', '')}\n")
                lines.append(f"```\n{p.get('full_image_prompt', '')}\n```\n")
            return "\n".join(lines)
        return ""

    def _export_docx(self, export_root, data, name):
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()
        style = doc.styles["Normal"]
        font = style.font
        font.name = "Calibri"
        font.size = Pt(11)

        proj = data.get("project", {})
        title = doc.add_heading(proj.get("name", name), level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        meta = doc.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta.add_run(f"Topic: {proj.get('topic', '')}\n").bold = True
        meta.add_run(f"Status: {proj.get('status', '')}\n")
        meta.add_run(f"Created: {proj.get('created', '')}\n")
        doc.add_paragraph("")

        research = data.get("research", {})
        if research.get("generated_research"):
            doc.add_heading("Research", level=1)
            doc.add_paragraph(research["generated_research"])

        script = data.get("script", {})
        if script.get("script_output"):
            doc.add_heading("Script", level=1)
            doc.add_paragraph(script["script_output"])

        storyboard = data.get("storyboard", {})
        if storyboard.get("scenes"):
            doc.add_heading("Storyboard", level=1)
            for scene in storyboard["scenes"]:
                doc.add_heading(f"Scene {scene.get('scene_number', '')}", level=2)
                doc.add_paragraph(f"Timestamp: {scene.get('timestamp', '')}")
                doc.add_paragraph(f"Narration: {scene.get('narration', '')}")
                doc.add_paragraph(f"Visual: {scene.get('visual_description', '')}")
                doc.add_paragraph(f"Camera: {scene.get('camera_direction', '')}")

        prompts = data.get("image_prompts", {})
        if prompts.get("prompts"):
            doc.add_heading("Image Prompts", level=1)
            for prompt in prompts["prompts"]:
                doc.add_heading(f"Scene {prompt.get('scene_number', '')}", level=2)
                p = doc.add_paragraph()
                run = p.add_run(prompt.get("full_image_prompt", ""))
                run.font.size = Pt(10)

        out = export_root / f"{name}.docx"
        doc.save(str(out))
        return out

    def _export_pdf(self, export_root, data, name):
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
        from reportlab.lib.enums import TA_CENTER

        out = export_root / f"{name}.pdf"
        doc = SimpleDocTemplate(str(out), pagesize=A4,
                                leftMargin=0.75 * inch, rightMargin=0.75 * inch,
                                topMargin=0.75 * inch, bottomMargin=0.75 * inch)

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle("Title2", parent=styles["Title"], alignment=TA_CENTER, fontSize=20)
        heading_style = ParagraphStyle("Heading2", parent=styles["Heading2"], fontSize=14, spaceBefore=12)
        body_style = styles["Normal"]
        meta_style = ParagraphStyle("Meta", parent=body_style, alignment=TA_CENTER, fontSize=10, textColor="#555555")

        story = []

        proj = data.get("project", {})
        story.append(Paragraph(proj.get("name", name), title_style))
        story.append(Spacer(1, 6))
        story.append(Paragraph(f"Topic: {proj.get('topic', '')}", meta_style))
        story.append(Paragraph(f"Status: {proj.get('status', '')}", meta_style))
        story.append(Paragraph(f"Created: {proj.get('created', '')}", meta_style))
        story.append(Spacer(1, 20))

        research = data.get("research", {})
        if research.get("generated_research"):
            story.append(Paragraph("Research", heading_style))
            for line in research["generated_research"].split("\n"):
                if line.strip():
                    story.append(Paragraph(line.strip(), body_style))
            story.append(Spacer(1, 12))

        script = data.get("script", {})
        if script.get("script_output"):
            story.append(Paragraph("Script", heading_style))
            for line in script["script_output"].split("\n"):
                if line.strip():
                    story.append(Paragraph(line.strip(), body_style))
            story.append(Spacer(1, 12))

        storyboard = data.get("storyboard", {})
        if storyboard.get("scenes"):
            story.append(Paragraph("Storyboard", heading_style))
            for scene in storyboard["scenes"]:
                story.append(Paragraph(f"Scene {scene.get('scene_number', '')}", styles["Heading3"]))
                story.append(Paragraph(f"Timestamp: {scene.get('timestamp', '')}", body_style))
                story.append(Paragraph(f"Narration: {scene.get('narration', '')}", body_style))
                story.append(Paragraph(f"Visual: {scene.get('visual_description', '')}", body_style))
                story.append(Paragraph(f"Camera: {scene.get('camera_direction', '')}", body_style))
                story.append(Spacer(1, 8))

        prompts = data.get("image_prompts", {})
        if prompts.get("prompts"):
            story.append(Paragraph("Image Prompts", heading_style))
            for prompt in prompts["prompts"]:
                story.append(Paragraph(f"Scene {prompt.get('scene_number', '')}", styles["Heading3"]))
                story.append(Paragraph(prompt.get("full_image_prompt", ""), body_style))
                story.append(Spacer(1, 8))

        doc.build(story)
        return out
